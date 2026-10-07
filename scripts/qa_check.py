#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Turkcell backlink yazısı teslim öncesi kontrolü.

Kullanım:
    python3 qa_check.py "Yazı.docx" [...]

Mekanik olarak doğrulanabilen brief kurallarını kontrol eder. Okuma gerektiren
üç kontrol (advertorial testi, tekrar testi, doğruluk testi) bu scriptin
kapsamı dışında; onlar SKILL.md bölüm 7'de.
"""
import re
import sys
import zipfile

from docx import Document

RAKIPLER = [
    "Vodafone", "Türk Telekom", "Turk Telekom", "TTNET", "TT Mobil",
    "Bimcell", "Pttcell", "Netgsm", "Vestel Mobil",
]
REKLAM = [
    "hemen başvur", "hemen alın", "kaçırmayın", "fırsatı yakala", "tıklayın",
    "en iyi", "en hızlı", "sorunsuz", "kusursuz", "avantajlı fiyat",
    "rakipsiz", "devrim niteliğinde", "tartışmasız", "mükemmel", "benzersiz", "en sevilen",
]
FIYAT = re.compile(r"₺|\bTL\b|\bindirim\b|\bkampanyalı fiyat\b|\baylık ücret\b", re.I)
KIYAS = ["tüm operatörlerde", "tüm operatörler için", "operatörlerde aynı", "farklılaşan taraf", "hangi kanaldan alınırsa alınsın"]
TOOLING = ["Claude", "MCP", "DataForSEO", "Screaming Frog", "Ahrefs", "SEOmonitor"]


def check(path):
    pasaj = "- Pasaj -" in path or "--pasaj" in sys.argv
    d = Document(path)
    paragraphs = [p for p in d.paragraphs if p.text.strip()]
    body = "\n".join(p.text for p in paragraphs)
    tables = "\n".join(
        c.text for t in d.tables for r in t.rows for c in r.cells
    )
    full = body + "\n" + tables
    low = full.lower()

    heads1 = [p.text for p in paragraphs if p.style.name == "Heading 1"]
    heads2 = [p.text for p in paragraphs if p.style.name == "Heading 2"]
    h1 = heads1[0] if heads1 else ""
    # H1'den sonraki ilk iki gövde paragrafı = giriş
    intro = " ".join(
        p.text for p in paragraphs if p.style.name not in ("Heading 1", "Heading 2")
    ).split("\n")[0]
    non_head = [p.text for p in paragraphs if p.style.name not in ("Heading 1", "Heading 2")]
    intro = " ".join(non_head[:2])

    words = len(full.split())
    brand = full.count("Turkcell")

    z = zipfile.ZipFile(path)
    rels = z.read("word/_rels/document.xml.rels").decode()
    urls = re.findall(r'Target="(http[^"]+)"', rels)
    xml = z.read("word/document.xml").decode()
    anchors = re.findall(
        r"<w:hyperlink[^>]*>.*?<w:t[^>]*>([^<]*)</w:t>", xml
    )

    rows = []

    def add(ok, label, detail=""):
        rows.append((ok, label, detail))

    add(600 <= words <= 1400, "Kelime sayısı 600-1.400", str(words))
    add("Turkcell" not in h1 and "Pasaj" not in h1, "H1'de marka yok", h1[:60])
    add("Turkcell" not in intro and "Pasaj" not in intro, "Giriş paragraflarında marka yok")
    if pasaj:
        add(2 <= brand <= 4, "Marka (Turkcell Pasaj) 2-4 kez geçiyor", str(brand))
        last_h2 = heads2[-1] if heads2 else ""
        add("Pasaj" in last_h2, "Son H2 Pasaj bölümü", last_h2[:60])
        add(len(heads2) >= 3, "En az 3 H2 var", str(len(heads2)))
        add(len(d.tables) >= 1, "Karşılaştırma tablosu var", "%d tablo öğesi" % len(d.tables))
    else:
        add(3 <= brand <= 6, "Marka 3-6 kez geçiyor", str(brand))
        add(len(heads2) >= 4, "En az 4 H2 var", str(len(heads2)))
        add(len(d.tables) >= 2, "Kutu/tablo var", "%d tablo öğesi" % len(d.tables))
    add("—" not in full, "Em dash yok")
    add("  " not in full, "Çift boşluk yok")
    hits = [w for w in RAKIPLER if w.lower() in low]
    add(not hits, "Rakip operatör adı yok", ", ".join(hits))
    hits = [w for w in REKLAM if w in low]
    add(not hits, "Reklam kalıbı yok", ", ".join(hits))
    add(not FIYAT.search(full), "Fiyat/kampanya ifadesi yok")
    add("hazır kart" not in low, "'hazır kart' kullanılmamış")
    hits = [w for w in KIYAS if w in low]
    add(not hits, "Operatör/mağaza kıyas cümlesi yok", ", ".join(hits))
    last_p = non_head[-1] if non_head else ""
    add("Turkcell" in last_p and re.search(r"(ebilir|abilir)siniz", last_p) is not None,
        "Son paragraf CTA cümlesi", last_p[:70])
    hits = [w for w in TOOLING if w in full]
    add(not hits, "Araç/otomasyon adı sızmamış", ", ".join(hits))
    add(bool(urls), "Link var", " · ".join(sorted(set(urls))))
    add(len(set(urls)) == len(urls), "Aynı URL'ye tek link", "%d link" % len(urls))
    add(bool(anchors) and all(len(a.split()) <= 6 for a in anchors),
        "Anchor metinleri kısa ve sabit", " · ".join(anchors))

    print("\n" + path + ("  [Pasaj kuralları]" if pasaj else ""))
    print("-" * 72)
    ok_all = True
    for ok, label, detail in rows:
        mark = "OK  " if ok else "EKSIK"
        ok_all &= ok
        line = "%-6s %s" % (mark, label)
        if detail:
            line += "  ->  %s" % detail
        print(line)
    print("-" * 72)
    print("SONUC: %s" % ("temiz" if ok_all else "duzeltme gerekiyor"))
    print("Elle bakilacaklar: advertorial testi · tekrar testi · dogruluk testi")
    return ok_all


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    results = [check(p) for p in sys.argv[1:] if not p.startswith("--")]
    sys.exit(0 if all(results) else 1)


if __name__ == "__main__":
    main()
