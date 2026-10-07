#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""JSON -> Turkcell backlink blog .docx

Kullanım:
    python3 build_docx.py article.json "Çıktı Adı.docx"

JSON şeması:
{
  "title": "Yazı Başlığı",
  "blocks": [
    {"type": "p",      "text": "Paragraf. [anchor](https://...) linki gömülebilir."},
    {"type": "h2",     "text": "Bölüm Başlığı"},
    {"type": "bullet", "lead": "Kalın giriş:", "text": "Açıklama cümlesi."},
    {"type": "box",    "label": "Not:", "text": "Çerçeveli uyarı metni."},
    {"type": "table",  "header": ["A","B"], "rows": [["a1","b1"]]}
  ]
}

Inline biçimlendirme: [metin](url) link, **metin** kalın.
Biçim, onaylı örnek yazılarla aynı: Arial 12 pt gövde, 14 pt kalın siyah başlık,
ortalanmış H1, çerçeveli tablolar, mavi altı çizili link.
"""
import json
import re
import sys

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Arial"
SIZE = Pt(12)
TOKEN = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*)")


def new_doc():
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for style in ("Heading 1", "Heading 2"):
        h = d.styles[style]
        h.font.name = FONT
        h.font.bold = True
        h.font.size = Pt(14)
        h.font.color.rgb = RGBColor(0, 0, 0)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    return d


def _plain_run(p, text, bold=False, italic=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = SIZE
    r.bold = bold
    r.italic = italic
    return r


def _link_run(p, text, url):
    rid = p.part.relate_to(
        url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    rpr.append(fonts)
    for tag, val in (("w:color", "1155CC"), ("w:sz", "24"), ("w:szCs", "24")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        rpr.append(e)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hl.append(run)
    p._p.append(hl)


def render(p, text, italic=False):
    """[metin](url) ve **kalın** işaretlemesini runs'a çevirir."""
    for chunk in TOKEN.split(text):
        if not chunk:
            continue
        m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", chunk)
        if m:
            _link_run(p, m.group(1), m.group(2))
            continue
        m = re.fullmatch(r"\*\*([^*]+)\*\*", chunk)
        if m:
            _plain_run(p, m.group(1), bold=True, italic=italic)
            continue
        _plain_run(p, chunk, italic=italic)


def add_paragraph(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    render(p, text)


def add_bullet(d, lead, text):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    if lead:
        _plain_run(p, lead, bold=True)
        text = " " + text.lstrip()
    render(p, text)


def add_box(d, label, text):
    t = d.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    p = t.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    if label:
        _plain_run(p, label, bold=True, italic=True)
        text = " " + text.lstrip()
    render(p, text, italic=True)
    d.add_paragraph()


def add_table(d, header, rows):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, cell_text in enumerate(header):
        _plain_run(t.rows[0].cells[i].paragraphs[0], cell_text, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            render(cells[i].paragraphs[0], value)
    d.add_paragraph()


def word_count(d):
    n = sum(len(p.text.split()) for p in d.paragraphs)
    n += sum(
        len(c.text.split()) for t in d.tables for r in t.rows for c in r.cells
    )
    return n


def build(spec, out_path):
    d = new_doc()
    h = d.add_heading(spec["title"], level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for block in spec["blocks"]:
        kind = block["type"]
        if kind == "p":
            add_paragraph(d, block["text"])
        elif kind == "h2":
            d.add_heading(block["text"], level=2)
        elif kind == "bullet":
            add_bullet(d, block.get("lead", ""), block["text"])
        elif kind == "box":
            add_box(d, block.get("label", "Not:"), block["text"])
        elif kind == "table":
            add_table(d, block["header"], block["rows"])
        else:
            raise ValueError("bilinmeyen blok tipi: %s" % kind)
    d.save(out_path)
    return word_count(d)


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        spec = json.load(f)
    words = build(spec, sys.argv[2])
    print("Yazildi: %s" % sys.argv[2])
    print("Kelime sayisi: %d" % words)
    if words < 600:
        print("UYARI: 600 kelimenin altinda. Konu yarim kalmis olabilir; eksik senaryoyu ekle, dolgu ekleme.")
    elif words > 1400:
        print("UYARI: 1.400 kelimenin uzerinde. Tekrar eden bolum var mi kontrol et.")


if __name__ == "__main__":
    main()
